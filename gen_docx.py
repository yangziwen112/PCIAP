# -*- coding: utf-8 -*-
"""生成一张示例图片，并创建一个带图片的 docx 文档"""
import os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
IMG_PATH = os.path.join(BASE, 'sample_image.png')
DOCX_PATH = os.path.join(BASE, 'sample_doc.docx')

# 1. 生成一张示例图片（渐变背景 + 文字）
W, H = 800, 500
img = Image.new('RGB', (W, H))
draw = ImageDraw.Draw(img)
# 简单渐变
for y in range(H):
    r = int(102 + (118 - 102) * y / H)
    g = int(126 + (75 - 126) * y / H)
    b = int(234 + (162 - 234) * y / H)
    draw.line([(0, y), (W, y)], fill=(r, g, b))

# 尝试加载中文字体，找不到则用默认
try:
    font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 48)
except Exception:
    font = ImageFont.load_default()

draw.text((W // 2 - 150, H // 2 - 30), '示例图片', fill=(255, 255, 255), font=font)
img.save(IMG_PATH)
print('图片已生成:', IMG_PATH)

# 2. 创建带图片的 docx 文档
doc = Document()

# 标题
title = doc.add_heading('示例文档', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 段落
p = doc.add_paragraph('这是一份由脚本自动生成的示例 Word 文档，包含文字和图片。')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 插入图片
doc.add_picture(IMG_PATH, width=Inches(5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# 图片说明
cap = doc.add_paragraph('图1：示例图片')
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in cap.runs:
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# 正文段落
doc.add_paragraph('这是正文第一段。该文档演示了如何通过 python-docx 生成包含图片的 Word 文档。')
doc.add_paragraph('这是正文第二段。你可以在此基础上扩展更多内容，例如表格、列表、样式等。')

doc.save(DOCX_PATH)
print('文档已生成:', DOCX_PATH)
