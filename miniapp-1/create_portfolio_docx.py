# -*- coding: utf-8 -*-
"""生成《杨子玟-微信小程序开发作品集》"""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '杨子玟-微信小程序开发作品集.docx'
ASSET_DIR = ROOT / 'docs' / 'portfolio_assets'
ASSET_DIR.mkdir(parents=True, exist_ok=True)

FONT = 'Microsoft YaHei'
FONT_PATH = 'C:/Windows/Fonts/msyh.ttc'

def font(size, bold=False):
    try:
        return ImageFont.truetype(FONT_PATH, size)
    except Exception:
        return ImageFont.load_default()

def draw_diagram(name, title, nodes, arrows, width=1500, height=850):
    path = ASSET_DIR / name
    img = Image.new('RGB', (width, height), '#F5F8FC')
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((35, 25, width-35, height-25), radius=28, fill='#FFFFFF', outline='#D8E2F0', width=4)
    d.text((70, 55), title, fill='#173B73', font=font(42, True))
    boxes = {}
    for node in nodes:
        x, y, w, h, label, color = node
        d.rounded_rectangle((x, y, x+w, y+h), radius=20, fill=color, outline='#B8C7DA', width=3)
        lines = label.split('\n')
        total = sum(42 for _ in lines)
        top = y + (h-total)//2
        for i, line in enumerate(lines):
            bbox = d.textbbox((0, 0), line, font=font(30, True))
            d.text((x+(w-(bbox[2]-bbox[0]))//2, top+i*42), line, fill='#17324D', font=font(30, True))
        boxes[label] = (x, y, w, h)
    for src, dst in arrows:
        x, y, w, h = boxes[src]
        x2, y2, w2, h2 = boxes[dst]
        start = (x+w, y+h//2) if x+w <= x2 else (x+w//2, y+h)
        end = (x2, y2+h2//2) if x+w <= x2 else (x2+w2//2, y2)
        d.line([start, end], fill='#5A78A5', width=6)
        # arrow head
        ex, ey = end
        d.polygon([(ex, ey), (ex-18, ey-10), (ex-18, ey+10)], fill='#5A78A5')
    img.save(path)
    return path

def make_assets():
    assets = {}
    assets['architecture'] = draw_diagram(
        '01-总体架构.png', '民大通总体架构',
        [(100, 180, 260, 120, '微信小程序\n页面层', '#E7F0FF'),
         (460, 180, 260, 120, 'API 云函数\n认证与业务', '#EAF8F0'),
         (820, 180, 260, 120, '云数据库\n公开数据投影', '#FFF5DF'),
         (1180, 180, 220, 120, '用户\n可见结果', '#F3EAFE'),
         (460, 480, 260, 120, 'Crawler\n真实来源采集', '#EAF8F0'),
         (820, 480, 260, 120, 'RAG / STAR\n检索与审核', '#FFE8E8')],
        [('微信小程序\n页面层','API 云函数\n认证与业务'), ('API 云函数\n认证与业务','云数据库\n公开数据投影'), ('云数据库\n公开数据投影','用户\n可见结果'), ('API 云函数\n认证与业务','Crawler\n真实来源采集'), ('API 云函数\n认证与业务','RAG / STAR\n检索与审核'), ('RAG / STAR\n检索与审核','用户\n可见结果')]
    )
    assets['star'] = draw_diagram(
        '02-STAR工作流.png', 'AI 助手 STAR 工作流',
        [(100, 200, 260, 120, 'S 情境识别\n清理问题与风险', '#E7F0FF'),
         (460, 200, 260, 120, 'T 意图路由\n选择工具与权限', '#EAF8F0'),
         (820, 200, 260, 120, 'A 执行动作\n时间、数据库、联网', '#FFF5DF'),
         (1180, 200, 220, 120, '生成回答\n结论优先', '#F3EAFE'),
         (820, 500, 260, 120, 'R 审核反思\n证据、隐私、长度', '#FFE8E8'),
         (1180, 500, 220, 120, '安全降级\n仍然给出可用答复', '#EEF2F7')],
        [('S 情境识别\n清理问题与风险','T 意图路由\n选择工具与权限'), ('T 意图路由\n选择工具与权限','A 执行动作\n时间、数据库、联网'), ('A 执行动作\n时间、数据库、联网','生成回答\n结论优先'), ('生成回答\n结论优先','R 审核反思\n证据、隐私、长度'), ('R 审核反思\n证据、隐私、长度','安全降级\n仍然给出可用答复')]
    )
    assets['roles'] = draw_diagram(
        '03-角色协作.png', '平台角色与责任边界',
        [(120, 180, 280, 140, '学生用户\n浏览、提问、发布', '#E7F0FF'),
         (610, 180, 280, 140, '管理员\n采集、审核、维护', '#FFF5DF'),
         (1100, 180, 280, 140, 'AI 助手\n理解、检索、解释', '#F3EAFE'),
         (360, 520, 280, 140, '来源采集器\n抓取、清洗、去重', '#EAF8F0'),
         (850, 520, 280, 140, '审核器\n拒绝幻觉与泄密', '#FFE8E8')],
        [('学生用户\n浏览、提问、发布','AI 助手\n理解、检索、解释'), ('管理员\n采集、审核、维护','来源采集器\n抓取、清洗、去重'), ('管理员\n采集、审核、维护','审核器\n拒绝幻觉与泄密'), ('来源采集器\n抓取、清洗、去重','AI 助手\n理解、检索、解释'), ('AI 助手\n理解、检索、解释','审核器\n拒绝幻觉与泄密')]
    )
    return assets

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color='173B73'):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True, 'FFFFFF')
        set_cell_shading(table.rows[0].cells[i], '2F5597')
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], 'F2F6FB')
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph('')
    return table

def add_image(doc, path, width=15, caption=''):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width))
    if caption:
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(4)
        for r in cp.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(100, 110, 125)

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)

def add_page(doc, title, subtitle=None, break_before=True):
    """添加章节标题；短章节可连续排版，避免页面底部产生大块留白。"""
    if break_before:
        doc.add_page_break()
    h = doc.add_heading(title, level=1)
    h.paragraph_format.space_before = Pt(8)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.style = 'Subtitle'

def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def configure(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 48, 66)
    for name, size, color in [('Title', 28, '173B73'), ('Heading 1', 19, '173B73'), ('Heading 2', 14, '2F5597'), ('Subtitle', 12, '60758C')]:
        st = styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
    section = doc.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    header = section.header.paragraphs[0]
    header.text = '杨子玟｜微信小程序开发作品集'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.name = FONT
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(120, 130, 145)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('民大通校园信息聚合平台  ·  个人开发作品集')
    run.font.name = FONT
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(120, 130, 145)

def build():
    assets = make_assets()
    doc = Document()
    configure(doc)

    # 1 cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    r = p.add_run('杨子玟')
    r.font.name = FONT; r.font.size = Pt(18); r.font.color.rgb = RGBColor(96, 117, 140)
    title = doc.add_heading('微信小程序开发作品集', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('从校园资讯到 AI 助手：中央民族大学校园信息化实践').alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_image(doc, ROOT / 'images' / '4ea780dd5bc1b40c8b0a4e71c1acc994.png', 6.8, '图 1  民大通首页产品截图')
    p = doc.add_paragraph('我以中央民族大学校园生活为真实场景，设计并实现了民大通校园信息聚合平台。')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(8)
    add_table(doc, ['项目类型', '核心场景', '技术关键词'], [[
        '微信小程序作品集', '校园资讯、考试竞赛、AI 问答、校园墙与二手交易', '云函数、云数据库、Crawler、RAG、STAR Agent'
    ]], [4.2, 8.2, 5.6])
    doc.add_paragraph('开发作品集  ·  2026 年 8 月').alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 2 abstract
    add_page(doc, '摘要与阅读说明', '我希望读者先理解问题，再理解技术。')
    add_para(doc, '我把校园里分散的通知、竞赛、考试、就业、学院动态和学生交流，组织成一个可持续更新的信息入口。项目的核心不是把网页简单搬进小程序，而是围绕“来源是否可靠、时间是否有效、与谁有关、下一步做什么”四个问题，建立从采集到行动的完整链路。')
    add_para(doc, '我在前端使用微信小程序承载轻量交互，在云函数中完成认证、内容治理、采集器运行和 AI 服务编排；在智能服务层使用 LangChain 与 LangGraph 搭建 STAR 工作流，使意图识别、工具调用、检索增强生成、回答审核和失败降级能够被独立测试。')
    add_para(doc, '这份作品集按“问题—产品—角色—数据—Agent—工程—验证—展望”的顺序展开。没有截图的页面，我使用实际页面结构、接口职责和状态变化进行完整说明。')
    add_table(doc, ['阅读章节', '我希望读者看到的重点'], [
        ('项目定位', '校园信息碎片化为什么值得用系统解决'),
        ('产品与角色', '学生、管理员、采集器、AI 助手和审核器如何分工'),
        ('技术逻辑', '小程序、API、数据库、Crawler、RAG 如何连通'),
        ('工程实践', 'GitHub 研究、测试、隐私、失败恢复和可迁移方法'),
    ], [3.5, 12.5])

    # 3 contents and motivation
    add_page(doc, '一、项目缘起：校园信息碎片化', '我从一个真实而具体的校园使用问题开始。')
    add_para(doc, '我第一次认真思考这个问题，是在寻找一条考试通知时：官网有正式文件，学院群里有同学转发，评论区又出现了另一种说法。我花了很久确认“哪一个是真的、时间有没有变、我到底要做什么”。那一刻我意识到，校园里缺的不是信息，而是一条能把信息讲明白、带到行动上的路径。')
    add_para(doc, '在中央民族大学，学校官网、学院网站、竞赛平台、公众号和群聊各自承担着不同职责。它们都很有价值，却很少以学生的视角组织在一起。相同的通知可能出现多个标题；报名开始、报名截止和考试日期常常被挤在一段话里；一条已经结束的消息还会在群里反复转发。')
    add_para(doc, '我把这种体验称为“校园信息碎片化”。民大通并不追求把所有网页搬进小程序，而是先替学习者回答四个小问题：这条消息从哪里来、现在还有效吗、和我有什么关系、我下一步该做什么。后续的采集、摘要、AI 和权限设计，都围绕这四个问题展开。')
    add_table(doc, ['传统信息体验', '民大通的处理方式', '用户得到的结果'], [
        ('多个网站分别查找', '来源目录统一采集', '从一个入口进入'),
        ('文章开头堆叠摘要', '抽取时间、地点、对象、行动项', '先看懂再决定是否打开原文'),
        ('旧日期继续被转发', '使用服务器时间判断有效节点', '避免把过期安排当成下一次'),
        ('模型自由发挥', '证据评分、回答审核、一次反思', '短而有依据的回答'),
    ], [4.5, 6.5, 5])

    # 4 overview
    add_page(doc, '二、产品定位与总体逻辑', '民大通是一套校园信息工作台，而不是资讯搬运页。')
    add_image(doc, assets['architecture'], 16, '图 2  民大通总体架构示意图')
    add_para(doc, '我将系统拆成四层。页面层负责让用户浏览、提问、发布和管理；业务层负责认证、权限、内容和消息；数据层负责存储公开资讯、用户行为、校园动态和 AI 会话；智能层负责采集、检索、时间判断、回答生成与审核。')
    add_para(doc, '这四层之间通过明确接口连接。页面不直接读取云数据库，只调用 API 云函数；AI 不读取完整数据库，只接收经过筛选的公开字段；管理员不直接修改采集器内部逻辑，而是从平台管理中心按业务分组运行任务。')
    add_bullets(doc, ['官方资讯与学生动态分开建模，避免把学校通知和个人经验混为同一权威等级。', '二手书和二手闲置使用数量、成色、位置、交易方式、原因等结构化规格，正文只承担补充说明。', '参考资料默认折叠，用户先看到结论；点击后才展开来源，降低阅读压力。', '所有无法核验的日期和地点都明确标注待确认，不用虚构内容填充页面。'])

    # 5 roles
    add_page(doc, '三、角色体系：每个角色只承担清晰责任', '我用角色边界控制复杂度，也控制数据暴露。')
    add_image(doc, assets['roles'], 16, '图 3  平台角色与责任边界')
    add_table(doc, ['角色', '主要职责', '不能越过的边界'], [
        ('学生用户', '浏览公开资讯、向助手提问、订阅、收藏、发布动态和二手物品', '不能查看他人私信、管理员数据和未公开个人记录'),
        ('游客用户', '浏览公开资讯，体验最多三次 AI 对话', '不能收藏、订阅、发布、查看历史和私信'),
        ('管理员', '维护来源、运行采集器、审核内容、发布和编辑官方资讯', '管理员入口只在登录后的个人中心出现'),
        ('来源采集器', '按来源发现页面、提取字段、去重、过滤并写入内容库', '采集失败时不能伪造成功数据'),
        ('AI 助手', '理解问题、选择工具、组织证据、生成面向用户的短回答', '不能读取密码、身份证、私信、收藏明细和管理员内部数据'),
        ('审核器', '检查证据、时间、隐私、长度和网页导航泄漏', '不因模型回答流畅就跳过来源与时效检查'),
    ], [3, 7, 7])
    add_para(doc, '我特别强调游客和登录用户的区别。游客不是“半个用户”，而是拥有明确公开访问范围的访问者；当用户点击收藏、订阅、发布或私信时，系统会说明为什么需要登录，并把登录后能获得的功能讲清楚。')

    # 6 UI pages
    add_page(doc, '四、页面体验：从首页到校园动态', '我把高频入口放在用户最容易理解的位置。')
    add_image(doc, ROOT / 'images' / '4ea780dd5bc1b40c8b0a4e71c1acc994.png', 6.5, '图 4  首页：分类资讯与近期事项入口')
    add_para(doc, '首页承担“第一次理解平台”的任务。我在首屏安排资讯分类、近期截止事项、搜索入口和 AI 助手入口，用户可以先浏览，也可以直接提问。首页卡片优先呈现来源、对象、时间和下一步，不把整篇文章压在第一屏。')
    add_image(doc, ROOT / 'images' / '29d57afee43e55e25a11bfd9bdcc1c7b.png', 6.5, '图 5  校园动态：官方资讯与学生内容分层')
    add_para(doc, '校园动态承担“真实交流”的任务。官方资讯显示来源和自动更新标识，学生帖子显示作者、分类和互动，二手书单独展示交易规格。动态页面自动刷新，但只把通过学生相关性筛选的官方内容放入聚合流。')
    add_bullets(doc, ['首页解决“我现在应该先看什么”。', '动态解决“学校通知之外，同学正在关注什么”。', '详情页解决“我需要核对原文和完整条件”。', '发布页解决“我如何用统一格式表达一件事”。'])

    # 7 other pages
    add_page(doc, '五、页面体验：AI、消息、个人中心与管理中心', '截图展示的是入口，真正的价值在于页面之间的连通。', break_before=False)
    add_image(doc, ROOT / 'images' / 'f7230eaa4df803e870f240146356a201.png', 6.5, '图 6  AI 助手：新建会话与折叠参考资料')
    add_para(doc, 'AI 助手支持新建对话、切换历史会话、文字提问、图片上传和游客三次体验。回答先显示结论，参考资料默认折叠；用户点击“查看参考资料”后，才展开最多三条来源。机器人头像使用轻量动态陪伴效果，但动画不会覆盖消息气泡和输入区域。')
    add_image(doc, ROOT / 'images' / 'ba28f63b5af1cd45b345d8150668a8ee.png', 6.5, '图 7  管理中心：按业务分组运行采集器')
    add_para(doc, '管理中心是运营入口。管理员可以分别更新教资、考研、四六级、民大主页、竞赛、三创赛、创新创业、挑战杯、就业实习和信息工程学院内容。每次运行都会返回扫描、写入、更新和过滤数量，失败时显示鉴权或来源问题，而不是静默显示空页面。')
    add_para(doc, '没有截图的消息页和个人中心也遵循同一套逻辑：消息页对游客显示登录引导，登录后显示会话列表；个人中心集中收藏、订阅、浏览历史、隐私设置和管理员入口。它们把个人数据放在可控边界内，不让 AI 或公开接口直接读取。')

    # 8 data model
    add_page(doc, '六、数据模型：让通知从长文章变成可行动事项', '我用字段表达事实，用正文承载细节。')
    add_table(doc, ['字段', '用途', '用户看到的表达'], [
        ('title', '标题和主题识别', '卡片标题'), ('sourceName / sourceUrl', '来源追溯', '官方来源与原文入口'),
        ('publishTime', '发布时间', '更新于某日'), ('registrationStartTime', '报名开始时间', '报名开始'),
        ('deadline', '报名或提交截止', '截止时间'), ('startTime / endTime', '考试、活动或赛程时间', '活动开始/考试时间'),
        ('location', '地点、校区或考场', '地点'), ('audience', '适用对象', '谁需要关注'),
        ('actionItem', '下一步行动', '建议马上做什么'), ('freshnessScore / evidenceScore', '时效与证据评价', '决定是否支撑回答'),
    ], [4, 7, 6])
    add_para(doc, '我把“报名开始”“报名截止”和“考试时间”拆成三个字段，是为了避免最常见的时间误答。AI 先读取服务器时间，再比较每个节点是否已经过去；如果本轮报名已结束但下一轮尚未公布，回答会明确区分两件事，不把考试日期冒充报名日期。')
    add_para(doc, '校园墙采用另一套结构。普通动态保存内容、标签、图片和互动；二手物品额外保存 marketDetails，包括 quantity、condition、location、tradeMethod 和 reason。发布时这些字段会自动同步到正文尾部，浏览者即使不打开详情也能看到交易关键信息。')

    # 9 crawler
    add_page(doc, '七、真实采集器：从来源发现到标准化发布', '我把爬虫设计成可恢复、可审计的来源适配器。', break_before=False)
    add_para(doc, '采集器位于 `cloudfunctions/crawler`，来源目录集中记录教师资格考试、考研、四六级、综合竞赛、三创赛、创新创业大赛、挑战杯、就业实习和信息工程学院等业务。管理员页面只调用 sourceGroup，不把每个站点的细节暴露给前端。')
    add_table(doc, ['阶段', '处理动作', '失败时的结果'], [
        ('发现', '访问来源页面，遵守超时、重定向和页面规则', '记录来源失败，不创建假内容'),
        ('提取', '识别标题、正文、时间、地点、对象和链接', '缺少关键字段则进入过滤统计'),
        ('标准化', '清理 HTML、导航词、重复段落和无关栏目', '保留原文链接，摘要不直接截取开头'),
        ('筛选', '按学生相关性、近七天窗口和业务关键词过滤', '返回 filtered 数量，方便管理员判断'),
        ('写入', '按 externalId 或 sourceUrl 更新，避免重复', '单来源失败不拖垮其他来源'),
        ('复核', '展示新增、更新、过滤和错误摘要', '管理员可重新运行单个业务组'),
    ], [3, 8, 6])
    add_para(doc, '我尤其关注民大主页动态的“学生相关性”。并不是学校网站上所有内容都适合进入校园动态；青苗计划、教资考场安排、竞赛报名、学院教学安排等与学生直接相关的内容优先保留，过于行政化、重复或缺乏行动价值的内容会被过滤。')

    # 10 STAR
    add_page(doc, '八、AI 助手：STAR 多智能体编排', '我用显式状态图替代把所有逻辑塞进一个提示词。')
    add_image(doc, assets['star'], 16, '图 8  STAR 工作流示意图')
    add_para(doc, 'STAR 是我对智能体流程的工程化拆分。S 负责识别问题处境和风险，T 负责决定意图与工具，A 负责执行数据库、时间和联网查询并生成答案，R 负责审核并在必要时反思一次。每一步都有输入、输出和可观察 trace，便于定位“是没有找到数据，还是回答没有通过审核”。')
    add_table(doc, ['节点', '输入', '输出'], [
        ('Situation', '用户问题、历史、图片和登录状态', '清理后的问题、风险标记、上下文'),
        ('Intent', '问题文本和情境', 'route、category、keywords、tools、requiresLogin'),
        ('Retrieval', '关键词、时间范围和分类', '证据列表、来源、时间字段和评分'),
        ('Tool', '路由决定的工具集合', 'current_time、public_database、web_search 结果'),
        ('Answer', '证据、工具结果、历史和图片', '短答案与折叠参考链接'),
        ('Reviewer', '草稿和质量规则', 'approved、score、feedback、一次重试或安全降级'),
    ], [3.5, 6.5, 7])

    # 11 agent details
    add_page(doc, '九、各智能体的作用与边界', '我让每个智能体只解决一种工程问题。', break_before=False)
    add_para(doc, '意图智能体先用规则识别高频问题，覆盖时间、报名、考场、竞赛、平台帮助、账户操作、闲聊和能力边界。规则优先的好处是稳定、快、可测试；复杂问题才交给模型进行补充理解。')
    add_para(doc, '检索智能体负责“找到可能相关的资料”，不直接决定最终答案。它会对标题、摘要、分类和时间节点评分，排除成绩通知、合格证明、网页导航等与问题不匹配的内容。工具智能体负责并行读取服务器时间和公开数据库，联网搜索只有在用户明确要求或本地资料不足时才启动。')
    add_para(doc, '回答智能体遵循“先结论，再关键点，最后行动”的表达原则。它不能输出模型状态、数据库原文、traceId 或内部故障；参考资料只作为佐证，不能代替正文回答。审核智能体会拒绝“相关信息：1、2、3”式资料清单、网页导航泄漏、无证据精确日期和隐私字段。')
    add_bullets(doc, ['时间问题必须调用 current_time，不能凭模型记忆回答“现在”。', '公共资讯只返回最小公开投影，不把用户私有数据混入 RAG。', '模型失败时使用本地规则生成直接答案，不返回乱码或内部错误。', '审核失败最多修订一次，仍不通过时进入安全降级，保证用户得到自然回应。'])

    # 12 auth privacy
    add_page(doc, '十、认证、权限与隐私设计', '我把“能不能看”放在“想不想看”之前。')
    add_para(doc, '游客可以浏览官方资讯、校园动态和公开帖子，并体验最多三次 AI 对话。游客尝试收藏、订阅、发布、评论、私信、查看历史或个人消息时，页面会用友好提示引导登录。登录用户通过账号与微信 OpenID 的绑定关系访问自己的数据；管理员接口只在管理员登录后的个人中心显示。')
    add_table(doc, ['数据类型', '游客', '登录用户', '管理员'], [
        ('官方公开资讯', '可浏览', '可浏览、收藏、订阅', '可维护'), ('校园动态', '可浏览', '可发布、评论、点赞', '可治理'),
        ('二手物品', '可浏览', '可发布和管理自己的内容', '可按平台规则治理'), ('AI 对话', '最多三次', '可持续使用并保存历史', '可查看服务状态，不查看用户私密正文'),
        ('收藏、订阅、历史、私信', '不可见', '仅本人可见', '不因管理员身份自动开放'),
    ], [5, 3.5, 4, 4.5])
    add_para(doc, '我在 API 云函数中使用绑定账号校验，不接受客户端单独传入一个 userId 就直接放行。AI 的公开数据库工具只接收公开字段；密码、身份证后四位、私信、收藏明细和管理员内部数据不会进入模型上下文。')

    # 13 database
    add_page(doc, '十一、数据库与接口逻辑', '我用集合职责拆分复杂业务，减少互相污染。', break_before=False)
    add_table(doc, ['集合', '保存内容', '主要使用页面'], [
        ('users', '账号、角色、绑定关系、设置', '登录、个人中心、权限'), ('contents', '官方资讯、来源、时间和摘要', '首页、详情、AI、管理中心'),
        ('sources / tags', '来源目录、栏目和订阅标签', '发布、订阅、采集器'), ('campus_posts', '校园墙和二手物品', '校园动态、发布、详情'),
        ('campus_comments / campus_likes', '评论、点赞及互动', '校园墙详情'), ('favorites / history', '收藏和浏览历史', '个人中心'),
        ('messages / friend_requests / friends', '私信与好友关系', '消息页、搜索页'), ('ai_messages / ai_usage / ai_runs', '会话、游客额度、运行审计', 'AI 助手、服务治理'),
        ('crawl_logs', '采集成功、过滤和失败信息', '管理中心'),
    ], [4, 8, 5])
    add_para(doc, '接口层通过 route 分发业务，例如 `feed/recommend`、`content/detail`、`campus/feed/list`、`rag/chat`、`crawler/run` 和 `ai/messages/list`。页面只关心稳定的业务结果，不直接依赖数据库内部结构；这样后续更换存储实现时，页面不需要全部重写。')

    # 14 admin and publishing
    add_page(doc, '十二、管理员运营与内容治理', '我把“自动采集”和“人工确认”放在同一条可追溯链路中。')
    add_image(doc, ROOT / 'images' / 'ba28f63b5af1cd45b345d8150668a8ee.png', 6.5, '图 9  管理中心产品截图')
    add_para(doc, '管理员打开平台管理中心后，可以先查看资讯、来源和标签数量，再进入采集任务。采集任务按照业务分组排列，运行中按钮显示状态；完成后展示扫描、写入、更新和过滤统计。管理员可以发布人工确认的校内通知，也可以编辑或删除不再有效的内容。')
    add_para(doc, '清理历史演示数据是独立维护操作，只删除标记为示例链接、随机图片、旧年份演示标题或演示校园动态的数据，不触碰真实采集和用户发布内容。这个设计让我能够在调试阶段清理旧数据，同时保护真实业务内容。')
    add_para(doc, '内容进入公开页面前至少经过三层判断：来源是否可访问，内容是否与学生相关，字段是否能支撑卡片和回答。没有真实链接或来源的内容不会为了“让首页不空”而被补进系统。')

    # 15 user journeys
    add_page(doc, '十三、典型使用流程', '我用场景验证页面、接口和智能体是否真正连通。', break_before=False)
    add_table(doc, ['场景', '用户动作', '系统处理', '最终结果'], [
        ('查询教资报名', '输入“教资什么时候报名”', '识别 upcoming，读取当前时间与公开数据库，过滤过期报名节点', '先给下一次有效报名或明确说明尚未公布'),
        ('寻找竞赛', '输入“信息工程学院最近有什么竞赛”', '识别学院与竞赛关键词，按近七天和学生相关性检索', '返回对象、时间、报名动作和来源'),
        ('发布二手书', '选择二手书并填写五项规格', '校验必填字段，图片上传，规格同步到正文', '形成可快速阅读和交易的帖子'),
        ('游客使用 AI', '连续提问三次', '按 OpenID 记录次数，达到上限后引导登录', '公开体验可用，个人数据仍受保护'),
        ('管理员更新民大主页', '点击“更新民大主页”', '调用 crawler 指定 sourceGroup，过滤无关内容', '统计新增、更新、过滤和失败原因'),
    ], [3.5, 5, 6, 3.5])
    add_para(doc, '我把场景测试作为产品验收的一部分。一个页面看起来正常，并不等于逻辑完整；只有从入口、接口、数据库、返回结果到下一步动作都能连通，用户才会感到系统真实可靠。')

    # 16 Git knowledge
    add_page(doc, '十四、GitHub 实践与我吸收的工程知识', '我不把 GitHub 当作代码展示墙，而把它作为学习和验证的工程记录。')
    add_para(doc, '我把 GitHub 当成一座公开的工程图书馆来阅读。有人把 LangGraph 的状态图画得很清楚，有人用课程项目展示 Agentic RAG 如何自校正，也有人把评测、追踪和失败恢复写成生产环境的基本功。我没有照搬某一个仓库，而是把这些项目共同强调的“可解释、可测试、可恢复”带回民大通。')
    add_para(doc, '这种学习方式改变了我的开发顺序。以前我会先想“模型能不能回答”，后来我会先问“问题应该路由到哪里、模型能看到哪些字段、证据不足时如何停下来、回答错误时谁负责修订”。因此，STAR 工作流不是为了让项目看起来复杂，而是为了让每一次回答都能找到来路。')
    add_table(doc, ['公开项目方向', '我吸收的方法', '在民大通中的落地'], [
        ('LangGraph / LangChain', '显式状态图、Runnable、工具边界', 'STAR 工作流、节点 trace 和一次修订'),
        ('Agent production practices', '护栏、评测、部署和可观测', 'reviewScore、retryCount、latencyMs、失败降级'),
        ('Agentic RAG courses', '相关性评分、自校正和分支路由', '报名、考场、成绩等问题的相关度排序'),
        ('Context engineering', '压缩历史、去重证据、减少 Token', '只保留最近相关消息和最多三条参考资料'),
        ('校园聚合项目', '轻量入口、活动聚合、用户视角', '将官方来源、学生动态和二手交易分层'),
    ], [5, 7, 6])
    add_para(doc, '我在 GitHub 公开仓库中只保留脱敏配置、产品截图、架构文档和可复现测试，不提交真实 AppID、API Key、Token、用户数据或视频缓存。提交前我会检查敏感信息、假链接和历史演示数据；每一条公开内容都要能说明来源、用途和删除方式。对我而言，作品集不仅要展示“做出来了什么”，也要让别人放心地阅读和复现。')
    add_para(doc, '项目仓库： https://github.com/yangziwen112/-ai')

    # 17 testing
    add_page(doc, '十五、质量评估与错误恢复', '我把“不能报错”拆成可执行的验证项目。')
    add_table(doc, ['检查层', '我验证的内容', '通过标准'], [
        ('语法检查', '全部 JavaScript 文件 node --check', '不出现语法错误'), ('小程序配置', 'app.json、项目页面和本地资源', '页面路径、图标和组件存在'),
        ('RAG 工作流', '情境、意图、检索、回答、审核和降级', 'STAR 测试全部通过'), ('采集器', '来源目录、分组和字段规则', '来源测试通过'),
        ('回答质量', '证据、简洁、行动性、日期支持、网页导航泄漏', '审核规则拒绝不合格答案'), ('权限安全', '游客、登录用户、管理员路由', '越权请求被拒绝'),
    ], [4, 8, 6])
    add_para(doc, '当模型服务不可用时，我不让前端出现“模型暂时繁忙”“数据库原文”或 RAG 内部错误，而是根据意图和证据生成自然的规则回答。回答审核失败时先修订一次；仍失败则安全降级，告诉用户目前能确认什么、还缺什么、下一步应如何核对。')
    add_bullets(doc, ['时间问题：没有未来节点时明确说尚未公布，不猜测日期。', '来源问题：没有可靠来源时不把搜索清单当答案。', '权限问题：需要登录时解释原因和收益，不只弹出无权限。', '采集问题：显示失败和恢复动作，不用旧示例填充页面。'])

    # 18 innovation
    add_page(doc, '十六、创新点与可迁移价值', '我希望这套方法不只服务一个学校。', break_before=False)
    add_table(doc, ['创新点', '我解决的问题', '未来可迁移方向'], [
        ('事项化资讯', '长通知无法快速判断行动', '课程、政务、社团和企业内部事项'), ('证据驱动 Agent', '模型可能复述过期或不相关内容', '法律、医疗之外的低风险垂直问答'),
        ('权限投影', '模型看到过多私有数据', '个人财务、学习档案和企业知识库'), ('来源适配器', '网站变化导致爬虫整体失效', '不同学校使用配置替换来源'),
        ('人机协同审核', '自动发布可能放大错误', '低证据内容进入人工复核队列'), ('个人校园记忆', '提醒过多或与用户无关', '用户授权后生成事项日程'),
    ], [4, 7, 7])
    add_para(doc, '我认为项目最有价值的创新不在于把多个模型堆在一起，而在于把校园信息的来源、时间、对象、权限和行动组织成可解释的系统。未来我会继续完善事项版本链、日历提醒、来源健康度、人工复核队列和跨学校适配能力。')
    add_para(doc, '如果把民大通迁移到另一所高校，我只需要替换来源目录、学校标签和内容筛选规则，核心的页面结构、权限模型、STAR 工作流、审核规则和测试方法仍然可以复用。这使它从一个校园小程序，逐步成为一套可迁移的校园信息化方法。')

    # 19 conclusion
    add_page(doc, '十七、总结：我交付的不只是一个小程序', '我把一次开发实践整理成一套可以继续生长的工程系统。', break_before=False)
    add_para(doc, '通过民大通，我完成了从产品定位、页面交互、云函数接口、数据库设计、真实采集、AI 编排、权限边界到测试验证的一整套闭环。教资只是其中一个业务入口，平台真正关注的是大学生每天会遇到的校园信息：考试考证、竞赛实践、学院通知、就业实习、活动讲座、校园交流和二手交易。')
    add_para(doc, '我坚持三个原则。第一，真实来源优先，不能用虚假示例替代真实业务；第二，回答必须直接解决问题，参考资料只作为证据；第三，系统必须有边界，能够明确拒绝、降级和引导。')
    add_para(doc, '这份作品集记录的是我的设计判断和工程选择。页面截图展示了用户看见的结果，架构和流程说明补齐了用户看不见的逻辑，GitHub 实践说明了我如何学习和吸收公开知识，测试部分说明了我如何确认改动没有破坏系统。')
    add_para(doc, '我希望这个项目能够让读者看到：一个真正有价值的校园智能入口，不是把信息和模型简单叠加，而是把问题拆清楚，把数据治理好，把权限守住，把答案变短，把下一步做实。')
    add_table(doc, ['项目成果', '对应文件或页面'], [
        ('产品页面与交互', 'pages/、components/、app.json'), ('业务 API 与权限', 'cloudfunctions/api'), ('真实来源采集', 'cloudfunctions/crawler'),
        ('STAR 智能体工作流', 'cloudfunctions/rag/lib/workflow.js'), ('可复用 Skill', 'skills/'), ('架构与设计资料', 'docs/architecture/、docs/'),
        ('公开作品仓库', 'https://github.com/yangziwen112/-ai'),
    ], [5.5, 12.5])
    doc.add_paragraph('')
    p = doc.add_paragraph('杨子玟')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)

    # 附录：把页面中不便完整展示的工程细节集中说明，避免正文出现大片留白
    add_page(doc, '附录 A、页面状态与交互检查表', '我用状态而不是静态截图描述页面是否真正可用。')
    add_table(doc, ['页面', '正常状态', '空状态', '异常与恢复'], [
        ('首页', '分类卡片、近期事项、搜索和 AI 入口', '显示真实来源为空的说明', '请求失败保留上次有效内容并提示刷新'),
        ('校园动态', '官方资讯、学生动态、二手物品分层', '显示发布入口和筛选条件', '分页失败可重试，不清空已有列表'),
        ('AI 助手', '结论优先、参考资料折叠、输入框固定', '显示示例问题和登录额度', '服务失败转为规则回答，不展示内部错误'),
        ('发布页', '标签、图片、文字和二手规格联动', '提示还可发布的内容类型', '上传失败保留文字草稿，避免重复填写'),
        ('管理中心', '业务分组、运行状态、统计与日志', '显示暂无任务结果', '鉴权失败单独提示，不伪造采集成功'),
    ], [3, 6, 4, 6])
    add_para(doc, '我在每个页面都区分加载中、空数据、请求失败和权限不足四种状态。这样用户不会把“暂时没有内容”误解成系统损坏，也不会因为一次网络波动丢失正在编辑的内容。')
    add_bullets(doc, ['按钮只绑定一个明确事件，避免发布、返回和卡片点击发生冒泡。', '列表使用稳定的 `_id` 或 `externalId`，避免更新后出现重复卡片。', '详情页返回时恢复原筛选条件和滚动位置，减少重复操作。', '所有需要登录的动作先解释用途，再提供登录按钮。'])

    add_page(doc, '附录 B、公开数据字段字典', '我用统一字段让不同来源可以被同一套页面和 Agent 使用。')
    add_table(doc, ['字段组', '字段', '校验规则', '使用位置'], [
        ('身份', 'sourceName、sourceUrl、externalId', '来源可追溯，externalId 稳定', '详情、引用、去重'),
        ('时间', 'publishTime、registrationStartTime、deadline、startTime、endTime', '使用 ISO 时间或可解析日期', '卡片、提醒、时间问答'),
        ('对象', 'category、audience、college、grade', '使用平台标签字典', '筛选、意图路由、推荐'),
        ('行动', 'actionItem、location、contact', '行动项必须是可执行短句', '摘要、回答、通知'),
        ('质量', 'freshnessScore、evidenceScore、relevanceScore', '0 到 1 之间，保留计算依据', '审核、排序、采集日志'),
        ('交易', 'quantity、condition、location、tradeMethod、reason', '二手物品发布必填', '校园墙、详情、搜索'),
    ], [4, 6, 5, 4])
    add_para(doc, '字段字典的价值在于把“网页文章”转换成“可操作事项”。当来源页面改版时，我只需要维护来源适配器，不需要同时修改首页、详情页和 AI 工作流。')
    add_para(doc, '对于缺失字段，我保留“未公布”或“待确认”，不使用推测值补全。所有时间回答都以当前服务器时间为基准，并优先选择未来最近的一次有效节点。')

    add_page(doc, '附录 C、发布前安全与质量清单', '我在提交代码和更新内容前，按清单完成最后一次检查。')
    add_table(doc, ['检查项', '具体动作', '不通过时的处理'], [
        ('敏感信息', '扫描 AppID、密钥、Token、Cookie、用户数据和本地缓存', '移除并改为环境变量占位提示'),
        ('来源真实性', '确认 URL 可访问、标题和正文对应、发布时间可解释', '进入过滤或人工复核，不写入公开流'),
        ('时间有效性', '用当前时间比较报名、截止、考试和活动节点', '标记已过期或选择下一次有效节点'),
        ('权限隔离', '用游客、普通用户、管理员三种身份测试接口', '拒绝越权并返回友好登录引导'),
        ('回答质量', '检查结论、证据、长度、行动项和隐私', '最多修订一次，仍失败则安全降级'),
        ('回归验证', '运行 RAG、crawler 测试并检查 JavaScript 语法', '定位失败节点，禁止带错发布'),
    ], [4, 8, 7])
    add_para(doc, '我把这份清单作为每次发布的最小门槛。它既适用于代码，也适用于采集内容和提示词调整，保证平台不会因为一次临时修改而重新出现假链接、旧日期、空首页或内部错误泄漏。')

    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()
